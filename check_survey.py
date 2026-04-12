import os

# Check if file exists
file_path = r"F:\MotiBuddy1\用户调查问卷结果\351800066_按序号_MotiBuddy 用户体验调查问卷_41_41.xlsx"

print(f"Checking file: {file_path}")
print(f"File exists: {os.path.exists(file_path)}")

if os.path.exists(file_path):
    print(f"File size: {os.path.getsize(file_path)} bytes")
    
    # Try to read with openpyxl
    try:
        from openpyxl import load_workbook
        wb = load_workbook(file_path)
        print(f"\n工作表数量: {len(wb.sheetnames)}")
        print(f"工作表名称: {wb.sheetnames}")
        
        # Get first sheet
        ws = wb.active
        print(f"\n活动工作表: {ws.title}")
        print(f"行数: {ws.max_row}")
        print(f"列数: {ws.max_column}")
        
        # Print headers
        headers = []
        for cell in ws[1]:
            headers.append(cell.value)
        print(f"\n表头: {headers}")
        
        # Print first few rows
        print("\n前3行数据:")
        for i, row in enumerate(ws.iter_rows(min_row=2, max_row=4, values_only=True), 2):
            print(f"Row {i}: {row}")
            
    except Exception as e:
        print(f"Error reading with openpyxl: {e}")
        
    # Try with pandas
    try:
        import pandas as pd
        df = pd.read_excel(file_path)
        print(f"\n=== Pandas读取成功 ===")
        print(f"记录数: {len(df)}")
        print(f"列名: {list(df.columns)}")
        print("\n前3行:")
        print(df.head(3))
    except Exception as e:
        print(f"Error reading with pandas: {e}")
