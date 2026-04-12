#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Convert Markdown thesis outline to Word document
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def convert_markdown_to_docx(md_file, docx_file):
    """Convert Markdown file to Word document"""
    
    # Read markdown content
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Create Word document
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Process markdown content line by line
    lines = content.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # Skip empty lines
        if not line:
            i += 1
            continue
        
        # Headers
        if line.startswith('# '):
            # Title (H1)
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(18)
                run.font.bold = True
        elif line.startswith('## '):
            # Chapter/Abstract (H2)
            doc.add_heading(line[3:], level=2)
        elif line.startswith('### '):
            # Section (H3)
            doc.add_heading(line[4:], level=3)
        elif line.startswith('#### '):
            # Subsection (H4)
            doc.add_heading(line[5:], level=4)
        elif line.startswith('**') and line.endswith('**'):
            # Bold text
            text = line.strip('*')
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.bold = True
        elif line.startswith('- ') or line.startswith('* '):
            # Bullet point
            text = line[2:]
            doc.add_paragraph(text, style='List Bullet')
        elif line[0].isdigit() and '. ' in line[:4]:
            # Numbered list
            text = line[line.find('. ')+2:]
            doc.add_paragraph(text, style='List Number')
        elif line.startswith('---'):
            # Horizontal rule - add page break
            doc.add_page_break()
        elif line.startswith('|') and '---' not in line:
            # Table row (simplified handling)
            if i + 1 < len(lines) and '---' in lines[i + 1]:
                # This is table header
                i += 2  # Skip header and separator
                continue
            # Add as regular text for now
            doc.add_paragraph(line)
        elif line.startswith('```'):
            # Code block
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i])
                i += 1
            code_text = '\n'.join(code_lines)
            p = doc.add_paragraph()
            run = p.add_run(code_text)
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
        else:
            # Regular paragraph
            if line:
                doc.add_paragraph(line)
        
        i += 1
    
    # Save document
    doc.save(docx_file)
    print(f"✅ Word document created successfully!")
    print(f"📄 Location: {docx_file}")
    return docx_file

if __name__ == '__main__':
    md_file = r"F:\MotiBuddy1\thesis_outline.md"
    docx_file = r"F:\MotiBuddy1\thesis_outline.docx"
    
    try:
        convert_markdown_to_docx(md_file, docx_file)
    except Exception as e:
        print(f"❌ Error: {e}")
        import traceback
        traceback.print_exc()
