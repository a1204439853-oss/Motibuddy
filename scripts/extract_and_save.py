from docx import Document
import os

docx_path = r'F:\MotiBuddy1\CUC海南02_本科毕设开题报告表-尹浩天.docx'
output_path = r'F:\MotiBuddy1\开题报告内容.txt'

if os.path.exists(docx_path):
    print('文件存在，开始提取...')
    doc = Document(docx_path)
    
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
    
    # 提取表格内容
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(' | '.join(row_text))
    
    result = '\n'.join(full_text)
    
    # 保存到文件
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(result)
    
    print(f'内容已保存到: {output_path}')
    print('---内容预览---')
    print(result[:2000])
else:
    print(f'文件不存在: {docx_path}')
