from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Read the markdown file
with open(r'F:\MotiBuddy1\thesis_outline.md', 'r', encoding='utf-8') as f:
    lines = f.readlines()

# Create a new Word document
doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# Process lines
for line in lines:
    line = line.rstrip()
    
    if not line:
        continue
    
    if line.startswith('# ') and not line.startswith('## '):
        # H1 - Title
        p = doc.add_heading(line[2:], level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith('## '):
        # H2
        doc.add_heading(line[3:], level=2)
    elif line.startswith('### '):
        # H3
        doc.add_heading(line[4:], level=3)
    elif line.startswith('#### '):
        # H4
        doc.add_heading(line[5:], level=4)
    elif line.startswith('- ') or line.startswith('* '):
        # Bullet
        doc.add_paragraph(line[2:], style='List Bullet')
    elif len(line) > 2 and line[0].isdigit() and '. ' in line[:4]:
        # Numbered list
        text = line[line.find('. ')+2:]
        doc.add_paragraph(text, style='List Number')
    elif line.startswith('---'):
        # Page break
        doc.add_page_break()
    elif line.startswith('|'):
        # Skip table lines for now
        continue
    elif line.startswith('```'):
        # Skip code block markers
        continue
    else:
        # Regular paragraph
        doc.add_paragraph(line)

# Save the document
output_path = r'F:\MotiBuddy1\thesis_outline.docx'
doc.save(output_path)
print(f"Document saved to: {output_path}")
